"""Exercise a built EXE: resources, spreadsheet editing, restart and shutdown.

Use --source to run the same API checks against editor.py on Linux/macOS.
"""
import argparse
import base64
import io
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
from urllib.parse import parse_qs, urlsplit
from urllib.error import HTTPError
from urllib.request import Request, urlopen

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'libraries'))
from openpyxl import Workbook, load_workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', action='store_true')
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='TableLens EXE test ', ignore_cleanup_errors=True) as temporary:
        folder = Path(temporary)
        appdata = folder / 'LocalAppData'
        log = appdata / 'TableLens' / 'TableLens.log'
        workspace = log.parent / 'workspace'
        log.parent.mkdir(parents=True)
        process = None
        output = None

        def start():
            nonlocal process, output
            log.write_text('')
            if args.source:
                output = log.open('w', encoding='utf-8')
                command = [sys.executable, '-B', str(ROOT / 'editor.py'), '--no-browser', '--storage-dir', str(workspace)]
            else:
                command = [str(ROOT / 'packages' / 'TableLens.exe'), '--no-browser']
            process = subprocess.Popen(command, cwd=folder, env={**os.environ, 'LOCALAPPDATA': str(appdata)},
                                       stdout=output or subprocess.DEVNULL, stderr=subprocess.STDOUT)
            deadline = time.monotonic() + 60
            while time.monotonic() < deadline:
                if process.poll() is not None:
                    raise RuntimeError(log.read_text(encoding='utf-8', errors='replace'))
                match = re.search(r'TableLens: (http://\S+)', log.read_text(encoding='utf-8', errors='replace'))
                if match:
                    return match.group(1)
                time.sleep(.1)
            raise RuntimeError('EXE did not start: ' + log.read_text(encoding='utf-8', errors='replace'))

        def api(address, action, data=None):
            token = parse_qs(urlsplit(address).query)['token'][0]
            url = address.split('/?')[0] + '/api/' + action
            request = Request(url, data=json.dumps(data or {}).encode(),
                              headers={'Content-Type': 'application/json', 'X-Editor-Token': token})
            try:
                with urlopen(request, timeout=20) as response:
                    return json.load(response)
            except HTTPError as error:
                raise RuntimeError(action + ': ' + error.read().decode('utf-8', errors='replace')) from error

        def stop(address):
            nonlocal process, output
            api(address, 'shutdown')
            assert process.wait(timeout=20) == 0
            process = None
            if output:
                output.close()
                output = None

        try:
            address = start()
            html = urlopen(address).read().decode()
            for resource in ['__TOKEN__', '__SYSTEM_LANGUAGE__', '__SYSTEM_THEME__']:
                assert resource not in html
            for element in ['theme-toggle', 'wizard-theme-toggle', 'language-toggle', 'wizard-language-toggle']:
                assert 'id="' + element + '"' in html
            for asset in ['pdf.mjs', 'pdf.worker.mjs']:
                url = address.replace('/?', '/assets/' + asset + '?')
                assert len(urlopen(url).read()) > 1000
            row = api(address, 'import', {'name': 'sample.tsv', 'content': base64.b64encode(b'Name\nFirst\n').decode()})
            values = list(row['values'])
            values[0] = 'Saved from EXE'
            api(address, 'save', {'revision': row['revision'], 'index': row['index'], 'values': values})
            document_id = row['document_id']
            assert b'Saved from EXE' in urlopen(address.replace('/?', '/download?')).read()
            book = Workbook()
            book.active.append(['Name'])
            book.active.append(['Excel row'])
            raw = io.BytesIO()
            book.save(raw)
            row = api(address, 'import', {'name': 'book.xlsx', 'content': base64.b64encode(raw.getvalue()).decode()})
            assert row['values'] == ['Excel row']
            api(address, 'save', {'revision': row['revision'], 'index': row['index'], 'values': ['Excel saved']})
            result = urlopen(address.replace('/?', '/download?')).read()
            assert load_workbook(io.BytesIO(result)).active['A2'].value == 'Excel saved'
            # Exercise SourcePreview and its bundled bridge without accessing the internet.
            row = api(address, 'import', {'name': 'source.tsv', 'content': base64.b64encode(
                ('Name\tSource\nFirst\t' + address + '\n').encode()).decode()})
            preview = api(address, 'source', {'index': row['index'], 'url': address})
            assert preview['kind'] == 'interactive'
            assert preview['bridge_key'].encode() in urlopen(preview['preview_url']).read()
            assert workspace.is_dir()
            stop(address)
            address = start()
            assert any(item['id'] == document_id for item in api(address, 'recent')['recent'])
            assert api(address, 'load', {'document_id': document_id})['values'][0] == 'Saved from EXE'
            stop(address)
            print('Resources, themes/languages, TSV/XLSX editing, source preview, persistent data and shutdown: OK')
        finally:
            if process is not None:
                # Shut down the app, including the child process in a one-file EXE.
                try:
                    api(address, 'shutdown')
                except Exception:
                    if sys.platform == 'win32' and not args.source:
                        subprocess.run(['taskkill', '/PID', str(process.pid), '/T', '/F'],
                                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                    else:
                        process.terminate()
                process.wait(timeout=20)
            if output:
                output.close()


if __name__ == '__main__':
    main()
