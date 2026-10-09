"""Table adapters. Originals are read only; serialization produces new bytes."""
import copy
import csv
import datetime as dt
import io
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

SUPPORTED = ('.tsv', '.csv', '.txt', '.xlsx', '.xlsm', '.ods', '.xls', '.xlsb')
TABLE = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
TEXT = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
OFFICE = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
XLINK = 'http://www.w3.org/1999/xlink'
MAX_ROWS, MAX_COLS = 100_000, 1_000


def display(value):
    if value is None:
        return ''
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    return str(value)


def scalar(new, old, formula=False):
    if isinstance(old, bool) and new.upper() in ('TRUE', 'FALSE'):
        return new.upper() == 'TRUE'
    if isinstance(old, (int, float)) and not isinstance(old, bool):
        try:
            return int(new) if re.fullmatch(r'[+-]?\d+', new) else float(new)
        except ValueError:
            pass
    if isinstance(old, (dt.datetime, dt.date, dt.time)):
        try:
            return type(old).fromisoformat(new)
        except ValueError:
            pass
    return new


class TableFile:
    def __init__(self, path, raw, options=None):
        self.path = Path(path)
        self.raw = raw
        self.options = dict(options or {})
        self.suffix = self.path.suffix.lower()
        if self.suffix not in SUPPORTED:
            raise ValueError('Поддерживаются TSV, CSV, TXT, XLSX, XLSM, ODS, XLS и XLSB.')
        self.has_header = self.options.get('header', True) is not False
        self.offset = 1 if self.has_header else 0
        self.links = {}
        self.sheets = []
        self.sheet = None
        self.workbook = None
        self.notice = ''
        if self.suffix in ('.tsv', '.csv', '.txt'):
            records = self.read_text()
        elif self.suffix in ('.xlsx', '.xlsm'):
            records = self.read_excel()
        elif self.suffix in ('.xls', '.xlsb'):
            records = self.read_legacy()
            self.notice = 'Результат будет XLSX. Данные всех листов сохраняются; оформление и макросы старого формата не переносятся.'
        else:
            records = self.read_ods()
        width = max((len(row) for row in records), default=0)
        if not width:
            raise ValueError('В выбранном листе нет данных.')
        if width > MAX_COLS or len(records) > MAX_ROWS:
            raise ValueError('Поддерживается до 100 000 строк и 1 000 колонок на листе.')
        self.headers = (records[0] + [''] * (width - len(records[0])) if self.has_header
                        else [f'Колонка {i+1}' for i in range(width)])
        self.rows = records[self.offset:]
        self.options.update({'header': self.has_header, 'sheet': self.sheet})
        self.output_suffix = '.xlsx' if self.suffix in ('.xls', '.xlsb') else self.suffix

    def read_text(self):
        encoding = self.options.get('encoding', 'auto')
        if encoding == 'auto':
            encoding = 'utf-16' if self.raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            try:
                content = self.raw.decode(encoding)
            except UnicodeDecodeError:
                encoding, content = 'cp1251', self.raw.decode('cp1251')
        else:
            try:
                content = self.raw.decode(encoding)
            except (UnicodeError, LookupError):
                raise ValueError('Не удалось прочитать файл. Выберите другую кодировку.')
        self.encoding = encoding
        self.bom = self.raw.startswith(b'\xef\xbb\xbf')
        self.line_end = '\r\n' if '\r\n' in content else '\n'
        delimiter = self.options.get('delimiter', 'auto')
        if delimiter == 'auto':
            if self.suffix == '.tsv':
                delimiter = '\t'
            else:
                try:
                    delimiter = csv.Sniffer().sniff(content[:65536], delimiters=',;\t|').delimiter
                except csv.Error:
                    delimiter = ',' if self.suffix == '.csv' else '\t'
        delimiter = '\t' if delimiter == 'tab' else delimiter
        if delimiter not in (',', ';', '\t', '|'):
            raise ValueError('Выберите разделитель: запятая, точка с запятой, табуляция или |.')
        self.delimiter = delimiter
        self.options.update({'delimiter': delimiter, 'encoding': encoding})
        return list(csv.reader(io.StringIO(content, newline=''), delimiter=delimiter))

    def select_sheet(self, names):
        self.sheets = names
        requested = self.options.get('sheet')
        self.sheet = requested if requested in names else names[0]

    def read_excel(self):
        from openpyxl import load_workbook
        self.workbook = load_workbook(io.BytesIO(self.raw), keep_vba=self.suffix == '.xlsm',
                                      data_only=False, keep_links=True, rich_text=True)
        self.select_sheet(self.workbook.sheetnames)
        return self.excel_records(self.workbook[self.sheet])

    def excel_records(self, sheet):
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLS:
            raise ValueError('Лист слишком большой: до 100 000 строк и 1 000 колонок.')
        records = []
        for number, cells in enumerate(sheet.iter_rows()):
            records.append([display(cell.value) for cell in cells])
            for column, cell in enumerate(cells):
                if cell.hyperlink and cell.hyperlink.target:
                    self.links[(number - self.offset, column)] = [cell.hyperlink.target]
        return records

    def read_legacy(self):
        from openpyxl import Workbook
        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        if self.suffix == '.xls':
            import xlrd
            source = xlrd.open_workbook(file_contents=self.raw)
            for sheet in source.sheets():
                target = self.workbook.create_sheet(sheet.name)
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLS:
                    raise ValueError('Лист слишком большой.')
                for row in sheet.get_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, source.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        values.append(value)
                    target.append(values)
                for (row, column), hyperlink in sheet.hyperlink_map.items():
                    if hyperlink.url_or_path:
                        target.cell(row+1, column+1).hyperlink = hyperlink.url_or_path
        else:
            from pyxlsb import open_workbook
            with open_workbook(io.BytesIO(self.raw)) as source:
                for name in source.sheets:
                    target = self.workbook.create_sheet(name)
                    with source.get_sheet(name) as sheet:
                        for row in sheet.rows():
                            for cell in row:
                                if cell.r >= MAX_ROWS or cell.c >= MAX_COLS:
                                    raise ValueError('Лист слишком большой.')
                                if cell.v is not None:
                                    target.cell(cell.r+1, cell.c+1, cell.v)
        self.select_sheet(self.workbook.sheetnames)
        return self.excel_records(self.workbook[self.sheet])

    def ods_text(self, cell):
        formula = cell.get(f'{{{TABLE}}}formula')
        if formula:
            return formula.split(':', 1)[-1]
        def walk(node):
            if node.tag == f'{{{TEXT}}}s':
                return ' ' * min(int(node.get(f'{{{TEXT}}}c', '1')), 10000)
            if node.tag == f'{{{TEXT}}}tab':
                return '\t'
            if node.tag == f'{{{TEXT}}}line-break':
                return '\n'
            result = node.text or ''
            for child in node:
                result += walk(child) + (child.tail or '')
            return result
        paragraphs = cell.findall(f'{{{TEXT}}}p')
        if paragraphs:
            return '\n'.join(walk(node) for node in paragraphs)
        return cell.get(f'{{{OFFICE}}}string-value', cell.get(f'{{{OFFICE}}}value',
                        cell.get(f'{{{OFFICE}}}date-value', cell.get(f'{{{OFFICE}}}boolean-value', ''))))

    def read_ods(self):
        with zipfile.ZipFile(io.BytesIO(self.raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200_000_000:
                raise ValueError('Распакованный файл слишком большой.')
            content = archive.read('content.xml')
        self.namespaces = {prefix: uri for _, (prefix, uri) in ET.iterparse(io.BytesIO(content), events=['start-ns'])}
        self.ods_root = ET.fromstring(content)
        self.ods_tables = self.ods_root.findall(f'.//{{{OFFICE}}}spreadsheet/{{{TABLE}}}table')
        names = [sheet.get(f'{{{TABLE}}}name', f'Лист {i+1}') for i, sheet in enumerate(self.ods_tables)]
        if not names:
            raise ValueError('В ODS нет листов.')
        self.select_sheet(names)
        self.ods_sheet = self.ods_tables[names.index(self.sheet)]
        self.ods_rows = []
        records = []
        for row in self.ods_sheet.iter(f'{{{TABLE}}}table-row'):
            values, links, spans = [], {}, []
            for cell in row:
                if cell.tag not in (f'{{{TABLE}}}table-cell', f'{{{TABLE}}}covered-table-cell'):
                    continue
                repeat = int(cell.get(f'{{{TABLE}}}number-columns-repeated', '1'))
                value = self.ods_text(cell)
                if not value and repeat > MAX_COLS:
                    repeat = max(0, MAX_COLS-len(values))
                if len(values)+repeat > MAX_COLS:
                    raise ValueError('В ODS слишком много колонок.')
                addresses = [node.get(f'{{{XLINK}}}href') for node in cell.iter(f'{{{TEXT}}}a')
                             if node.get(f'{{{XLINK}}}href')]
                for _ in range(repeat):
                    if addresses:
                        links[len(values)] = addresses
                    values.append(value)
                    spans.append(cell)
            while values and not values[-1]:
                values.pop()
            repeat = int(row.get(f'{{{TABLE}}}number-rows-repeated', '1'))
            if not values and repeat > MAX_ROWS:
                continue
            if len(records)+repeat > MAX_ROWS:
                raise ValueError('В ODS слишком много строк.')
            for _ in range(repeat):
                number = len(records)
                records.append(values[:])
                self.ods_rows.append((row, spans))
                for column, targets in links.items():
                    self.links[(number-self.offset, column)] = targets
        return records

    def serialize(self, rows):
        if self.suffix in ('.tsv', '.csv', '.txt'):
            output = io.StringIO(newline='')
            writer = csv.writer(output, delimiter=self.delimiter, lineterminator=self.line_end)
            if self.has_header:
                writer.writerow(self.headers)
            writer.writerows(rows)
            encoding = self.encoding
            if encoding == 'utf-8-sig' and not self.bom:
                encoding = 'utf-8'
            try:
                return output.getvalue().encode(encoding)
            except UnicodeError:
                raise ValueError('Новый текст не помещается в исходную кодировку. Откройте UTF-8 копию файла.')
        if self.suffix == '.ods':
            return self.serialize_ods(rows)
        # Work on a copy so a failed disk write cannot corrupt the loaded model.
        if self.suffix in ('.xlsx', '.xlsm'):
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(self.raw), keep_vba=self.suffix=='.xlsm',
                                    data_only=False, keep_links=True, rich_text=True)
        else:
            workbook = copy.deepcopy(self.workbook)
        sheet = workbook[self.sheet]
        for index, row in enumerate(rows):
            old = self.rows[index]
            for column, new in enumerate(row):
                if new == (old[column] if column < len(old) else ''):
                    continue
                cell = sheet.cell(index+self.offset+1, column+1)
                if type(cell).__name__ == 'MergedCell':
                    raise ValueError('Объединённую ячейку нужно редактировать через её верхнюю левую ячейку.')
                was_formula = cell.data_type == 'f'
                cell.value = scalar(new, cell.value)
                if isinstance(cell.value, str) and cell.value.startswith('=') and not was_formula:
                    cell.data_type = 's'
                if cell.hyperlink and new.startswith(('https://', 'http://')):
                    cell.hyperlink.target = new
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def serialize_ods(self, rows):
        root = copy.deepcopy(self.ods_root)
        sheets = root.findall(f'.//{{{OFFICE}}}spreadsheet/{{{TABLE}}}table')
        sheet = sheets[self.sheets.index(self.sheet)]
        # Expand repeated rows/cells only in the selected sheet; retain other sheets.
        parents = {child: parent for parent in sheet.iter() for child in parent}
        expanded = []
        for row in list(sheet.iter(f'{{{TABLE}}}table-row')):
            repeat = int(row.get(f'{{{TABLE}}}number-rows-repeated', '1'))
            if len(expanded) >= len(rows)+self.offset:
                break
            parent = parents[row]
            position = list(parent).index(row)
            amount = min(repeat, len(rows)+self.offset-len(expanded))
            parent.remove(row)
            for number in range(amount):
                replacement = copy.deepcopy(row)
                replacement.attrib.pop(f'{{{TABLE}}}number-rows-repeated', None)
                parent.insert(position+number, replacement)
                expanded.append(replacement)
            if repeat > amount:
                remainder = copy.deepcopy(row)
                remainder.set(f'{{{TABLE}}}number-rows-repeated', str(repeat-amount))
                parent.insert(position+amount, remainder)
        for index, values in enumerate(rows):
            old = self.rows[index]
            changes = {col: value for col, value in enumerate(values)
                       if value != (old[col] if col < len(old) else '')}
            if not changes:
                continue
            row = expanded[index+self.offset]
            cells = []
            for cell in list(row):
                if cell.tag not in (f'{{{TABLE}}}table-cell', f'{{{TABLE}}}covered-table-cell'):
                    continue
                repeat = int(cell.get(f'{{{TABLE}}}number-columns-repeated', '1'))
                position = list(row).index(cell)
                row.remove(cell)
                amount = min(repeat, max(changes)+1-len(cells))
                for number in range(max(0, amount)):
                    replacement = copy.deepcopy(cell)
                    replacement.attrib.pop(f'{{{TABLE}}}number-columns-repeated', None)
                    row.insert(position+number, replacement)
                    cells.append(replacement)
                if repeat > amount:
                    remainder = copy.deepcopy(cell)
                    remainder.set(f'{{{TABLE}}}number-columns-repeated', str(repeat-max(amount,0)))
                    row.insert(position+max(amount,0), remainder)
                if len(cells) > max(changes):
                    break
            while len(cells) <= max(changes):
                cell = ET.SubElement(row, f'{{{TABLE}}}table-cell')
                cells.append(cell)
            for column, value in changes.items():
                cell = cells[column]
                if cell.tag.endswith('covered-table-cell'):
                    raise ValueError('Нельзя редактировать закрытую часть объединённой ячейки ODS.')
                original_formula = cell.get(f'{{{TABLE}}}formula')
                original_type = cell.get(f'{{{OFFICE}}}value-type', 'string')
                for key in list(cell.attrib):
                    if key.startswith(f'{{{OFFICE}}}') or key == f'{{{TABLE}}}formula':
                        del cell.attrib[key]
                for child in list(cell):
                    if child.tag == f'{{{TEXT}}}p':
                        cell.remove(child)
                cell.set(f'{{{OFFICE}}}value-type', 'string')
                if original_formula and value.startswith('='):
                    cell.set(f'{{{TABLE}}}formula', original_formula.split(':',1)[0]+':'+value)
                elif original_type in ('float', 'percentage', 'currency'):
                    try:
                        number = float(value)
                        cell.set(f'{{{OFFICE}}}value-type', original_type)
                        cell.set(f'{{{OFFICE}}}value', str(number))
                    except ValueError:
                        pass
                elif original_type == 'boolean' and value.lower() in ('true','false'):
                    cell.set(f'{{{OFFICE}}}value-type', 'boolean')
                    cell.set(f'{{{OFFICE}}}boolean-value', value.lower())
                for paragraph in value.split('\n'):
                    p = ET.SubElement(cell, f'{{{TEXT}}}p')
                    # Explicit ODF spaces keep leading and repeated whitespace.
                    start = 0
                    for match in re.finditer(r' +|\t', paragraph):
                        segment = paragraph[start:match.start()]
                        if len(p):
                            p[-1].tail = segment
                        else:
                            p.text = segment
                        node = ET.SubElement(p, f'{{{TEXT}}}tab' if match[0]=='\t' else f'{{{TEXT}}}s')
                        if match[0]!='\t':
                            node.set(f'{{{TEXT}}}c', str(len(match[0])))
                        start = match.end()
                    if len(p):
                        p[-1].tail = paragraph[start:]
                    else:
                        p.text = paragraph
        for prefix, uri in self.namespaces.items():
            if not re.fullmatch(r'ns\d+', prefix):
                ET.register_namespace(prefix, uri)
        content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
        # Formula prefixes can exist only in attribute values; retain declarations.
        for prefix, uri in self.namespaces.items():
            name = ('xmlns:'+prefix if prefix else 'xmlns').encode()
            if name+b'=' not in content:
                root.set(name.decode(), uri)
        content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(self.raw)) as source, zipfile.ZipFile(output, 'w') as target:
            for item in source.infolist():
                target.writestr(item, content if item.filename=='content.xml' else source.read(item.filename))
        return output.getvalue()
